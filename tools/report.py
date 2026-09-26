"""Turn a graded session (result.json) into sheet-ready output.

- Computes the suggested penalty from config/rules.json (deterministic, not AI).
- Adds a history hint per student (prior sessions only) so the teacher can
  soften or harden the penalty.
- Writes cham_bai_AI.xlsx (same layout and colors as the Google Sheet) and
  paste.tsv (select-all, copy, paste into Google Sheets).
- If the teacher already graded this session in the Sheet, writes compare.md.

Usage:
    python tools/report.py TA6 19/9
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parent))
import history  # noqa: E402

PROJECT = Path(__file__).resolve().parent.parent
RED = "FFE02020"
PINK, YELLOW, ORANGE = "FFF0C8F8", "FFFFFF00", "FFFFE5CC"


def day_month(date: str) -> tuple[int, int]:
    d, m = re.findall(r"\d+", date)[:2]
    return int(d), int(m)


def session_dir(class_name: str, date: str) -> Path:
    d, m = day_month(date)
    return PROJECT / "data" / class_name.upper() / f"{d}-{m}"


def load_rules(class_name: str) -> dict:
    cfg = json.loads((PROJECT / "config" / "rules.json").read_text(encoding="utf-8"))
    rules = dict(cfg["default"])
    rules.update(cfg.get("classes", {}).get(class_name.upper(), {}))
    return rules


def pick_level(levels: list[dict], wrong: int, total: int) -> str:
    hit = [lv for lv in levels if wrong >= lv["min_wrong"]]
    return hit[-1]["penalty"].format(total=total) if hit else ""


def suggest_penalty(student: dict, rules: dict) -> str:
    if "penalty_override" in student:
        return student["penalty_override"]
    parts = student.get("parts", {})
    wrong = sum(len(p["wrong"]) for col, p in parts.items() if col in rules["counted_columns"])
    total = sum(p["total"] for col, p in parts.items() if col in rules["counted_columns"])
    out = [pick_level(rules["levels"], wrong, total)]
    for col, levels in rules.get("separate_columns", {}).items():
        if col in parts:
            out.append(pick_level(levels, len(parts[col]["wrong"]), parts[col]["total"]))
    return "; ".join(p for p in out if p)


def history_hint(name: str, prior: list[dict]) -> str:
    rows = [s["students"][name] for s in prior if name in s["students"]]
    if not rows:
        return ""
    penalties = sum(any("PHẠT" in k for k in r) for r in rows)
    scores = [history.score(r.get("TỪ VỰNG")) for r in rows]
    scores = [s for s in scores if s]
    rate = sum(a for a, _ in scores) / sum(b for _, b in scores) if scores else None
    if penalties == 0 and rate is not None and rate >= 0.9:
        return f"Lịch sử tốt ({len(rows)} buổi trước không bị phạt, đúng {rate:.0%}) - có thể giảm nhẹ"
    if penalties >= 2:
        return f"Đã bị phạt {penalties}/{len(rows)} buổi trước"
    return ""


def build_rows(result: dict, rules: dict, prior: list[dict]) -> list[dict]:
    rows = []
    for st in result["students"]:
        row = {"name": st["name"], "cells": {}, "penalty": "", "red_score": False, "alt": "",
               "notes": list(st.get("notes", []))}
        if st.get("status"):
            row["notes"].insert(0, st["status"])
        else:
            wrong_all = []
            for part in result["parts"]:
                p = st["parts"].get(part["column"])
                if p:
                    row["cells"][part["column"]] = f"{p['correct']}/{p['total']} {part['unit']}"
                    wrong_all += p["wrong"]
            row["cells"]["TỪ VIẾT SAI"] = ", ".join(wrong_all)
            row["penalty"] = suggest_penalty(st, rules)
            vocab = st["parts"].get("TỪ VỰNG")
            row["red_score"] = bool(vocab and len(vocab["wrong"]) >= 5)
            if vocab and vocab.get("meaning_wrong"):
                # score if Vietnamese meanings were not checked, for comparison only
                row["alt"] = f"{vocab['correct'] + len(vocab['meaning_wrong'])}/{vocab['total']} từ"
            if row["penalty"]:
                hint = history_hint(st["name"], prior)
                if hint:
                    row["notes"].append(hint)
        rows.append(row)
    return rows


def columns(result: dict) -> list[str]:
    cols = [p["column"] for p in result["parts"]]
    return cols[:1] + ["TỪ VIẾT SAI"] + cols[1:] + ["CHÉP PHẠT"]


def write_xlsx(result: dict, rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{result['class']} {result['date'].replace('/', '-')}"
    cols = columns(result)
    thin = Side(style="thin", color="FFBBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True, size=12)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.cell(1, 2, result.get("title", "TỪ VỰNG"))
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=len(cols))
    ws.cell(1, 2).font, ws.cell(1, 2).alignment = bold, center
    ws.cell(1, 2).fill = PatternFill("solid", fgColor=PINK)

    headers = [result["date"]] + cols + ["GHI CHÚ AI (xóa trước khi dán)"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(2, c, h)
        cell.font, cell.alignment, cell.border = bold, center, border
        is_yellow = c == 1 or h == "CHÉP PHẠT"
        cell.fill = PatternFill("solid", fgColor=YELLOW if is_yellow else PINK if c <= len(cols) + 1 else ORANGE)

    for r, row in enumerate(rows, 3):
        ws.cell(r, 1, row["name"]).font = bold
        for c, col in enumerate(cols, 2):
            value = row["penalty"] if col == "CHÉP PHẠT" else row["cells"].get(col, "")
            cell = ws.cell(r, c, value)
            cell.alignment = center
            red = (col == "TỪ VIẾT SAI" and row["penalty"]) or (col == "TỪ VỰNG" and row["red_score"])
            cell.font = Font(size=12, color=RED if red else None)
        note = ws.cell(r, len(cols) + 2, "\n".join(row["notes"]))
        note.alignment = Alignment(wrap_text=True, vertical="top")
        if row["notes"]:
            note.fill = PatternFill("solid", fgColor=ORANGE)
        for c in range(1, len(cols) + 3):
            ws.cell(r, c).border = border

    ws.column_dimensions["A"].width = 18
    for c in range(2, len(cols) + 2):
        ws.column_dimensions[ws.cell(2, c).column_letter].width = 24
    ws.column_dimensions[ws.cell(2, len(cols) + 2).column_letter].width = 70
    wb.save(path)


def write_tsv(result: dict, rows: list[dict], path: Path) -> None:
    cols = columns(result)
    lines = ["\t".join([result["date"]] + cols)]
    for row in rows:
        vals = [row["penalty"] if c == "CHÉP PHẠT" else row["cells"].get(c, "") for c in cols]
        lines.append("\t".join([row["name"]] + vals))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def compare(result: dict, rows: list[dict], teacher: dict) -> str:
    has_alt = any(row["alt"] for row in rows)
    head = "| Học sinh | AI | " + ("AI nếu bỏ qua nghĩa | " if has_alt else "") + "Cô | AI phạt | Cô phạt | Điểm khớp |"
    out = [f"# So sánh AI với cô - {result['class']} {result['date']}", "",
           head, "|" + "---|" * (head.count("|") - 1)]
    same = same_alt = graded = 0
    for row in rows:
        t = teacher["students"].get(row["name"], {})
        ai, tv = row["cells"].get("TỪ VỰNG", ""), t.get("TỪ VỰNG", "")
        t_pen = next((v for k, v in t.items() if "PHẠT" in k), "")
        mark = ""
        if history.score(ai) and history.score(tv):
            graded += 1
            ok = history.score(ai) == history.score(tv)
            ok_alt = history.score(row["alt"] or ai) == history.score(tv)
            same, same_alt = same + ok, same_alt + ok_alt
            mark = "khớp" if ok else ("khớp nếu bỏ qua nghĩa" if ok_alt else "LỆCH")
        alt = f"{row['alt'] or '-'} | " if has_alt else ""
        out.append(f"| {row['name']} | {ai or '-'} | {alt}{tv or '-'} | {row['penalty'] or '-'} | {t_pen or '-'} | {mark} |")
    out += ["", f"Điểm khớp: {same}/{graded} bài có điểm ở cả hai bên."]
    if has_alt:
        out.append(f"Nếu bỏ qua lỗi nghĩa tiếng Việt: {same_alt}/{graded}.")
    unmatched = result.get("unmatched_photos", [])
    if unmatched:
        out += ["", "Ảnh không ghép được với học sinh nào:"]
        out += [f"- {u['photo']} ({u['written_name']}): {u['note']}" for u in unmatched]
    return "\n".join(out)


def main(class_name: str, date: str) -> None:
    folder = session_dir(class_name, date)
    result = json.loads((folder / "result.json").read_text(encoding="utf-8"))
    hist = history.parse_class(class_name)
    this = day_month(date)
    prior = [s for s in hist["sessions"] if s["date"] and day_month(s["date"])[::-1] < this[::-1]]
    rows = build_rows(result, load_rules(class_name), prior)

    write_xlsx(result, rows, folder / "cham_bai_AI.xlsx")
    write_tsv(result, rows, folder / "paste.tsv")
    print(f"Wrote {folder / 'cham_bai_AI.xlsx'}\nWrote {folder / 'paste.tsv'}")

    teacher = next((s for s in hist["sessions"] if s["date"] and day_month(s["date"]) == this), None)
    if teacher:
        report = compare(result, rows, teacher)
        (folder / "compare.md").write_text(report + "\n", encoding="utf-8")
        print(report)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
