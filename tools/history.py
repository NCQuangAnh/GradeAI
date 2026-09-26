"""Parse past grading results from the Google Sheet export (data/sheet.xlsx).

Each tab is one class ("Tiếng anh 6" = TA6). Inside a tab, every session is a
block that starts with a date in column A, followed by a header row and one
row per student.

Usage:
    python tools/history.py TA6            # roster + per-student history summary
    python tools/history.py TA6 --json     # full history as JSON
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

SHEET = Path(__file__).resolve().parent.parent / "data" / "sheet.xlsx"


def tab_for(class_name: str, wb) -> str:
    code = re.sub(r"^TA", "", class_name.upper())
    for name in wb.sheetnames:
        if name.lower().replace("tiếng anh", "").strip() == code.lower():
            return name
    raise SystemExit(f"No tab for {class_name}. Tabs: {wb.sheetnames}")


def parse_class(class_name: str) -> dict:
    wb = openpyxl.load_workbook(SHEET, data_only=True)
    ws = wb[tab_for(class_name, wb)]
    sessions, roster = [], []
    date, headers = None, None
    for row in ws.iter_rows(values_only=True):
        a, rest = row[0], row[1:]
        filled = {i + 1: str(v).strip() for i, v in enumerate(rest) if v not in (None, "")}
        if isinstance(a, dt.datetime):
            date, headers = a.strftime("%d/%m"), None
        if a is None or isinstance(a, dt.datetime):
            cols = {i: t for i, t in filled.items() if not t.lower().startswith(("luật", "done"))}
            if len(cols) >= 2:
                headers = {i: t.upper() for i, t in cols.items()}
                sessions.append({"date": date, "columns": list(headers.values()), "students": {}})
            continue
        if isinstance(a, str) and a.strip() and headers:
            name = a.strip()
            if name not in roster:
                roster.append(name)
            sessions[-1]["students"][name] = {headers[i]: filled[i] for i in headers if i in filled}
    return {"class": class_name.upper(), "tab": ws.title, "roster": roster, "sessions": sessions}


def score(cell: str | None) -> tuple[int, int] | None:
    m = re.match(r"\s*(\d+)\s*/\s*(\d+)", cell or "")
    return (int(m.group(1)), int(m.group(2))) if m else None


def summary(data: dict) -> str:
    lines = [f"{data['class']} ({data['tab']}) - {len(data['sessions'])} buổi: "
             + ", ".join(s["date"] for s in data["sessions"])]
    for name in data["roster"]:
        parts, penalties, right, total = [], 0, 0, 0
        for s in data["sessions"]:
            r = s["students"].get(name)
            if not r:
                parts.append(f"{s['date']}: -")
                continue
            vocab = next((v for k, v in r.items() if k == "TỪ VỰNG"), None)
            pen = next((v for k, v in r.items() if "PHẠT" in k), None)
            sc = score(vocab)
            if sc:
                right, total = right + sc[0], total + sc[1]
            penalties += bool(pen)
            parts.append(f"{s['date']}: {vocab or '?'}" + (f" [phạt: {pen}]" if pen else ""))
        rate = f"{100 * right / total:.0f}%" if total else "n/a"
        lines.append(f"- {name}: đúng {rate}, bị phạt {penalties} lần | " + "; ".join(parts))
    return "\n".join(lines)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    data = parse_class(sys.argv[1])
    print(json.dumps(data, ensure_ascii=False, indent=2) if "--json" in sys.argv else summary(data))
